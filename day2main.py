import openpyxl
from openpyxl.chart import BarChart, Reference

myDataList=[]

def input_data():
    print("1.입력")
    input_continue = 'y'

    while input_continue != 'n':
        name = input("이름:")
        korean = int(input("국어:"))
        english = int(input("영어:"))
        myDataList.append({'name':name, 'korean':korean, 'english':english})
        
        input_continue = input("계속 입력하시겠습니까(y/n)?")
        if input_continue =='y':
            continue
        else:
            break    
        
    
def print_data():
    print("2.출력")
    #print("--------------------------")
    print("-"*30)
    print(" 이름           국어           영어      ")
    #print("--------------------------"    )
    print("-"*30)
    for n in myDataList:
        print("%(name)5s %(korean)5d %(english)5d" % n ) 
    
def save_excel_data():
    count = 0
    wb=openpyxl.Workbook()
    ws1=wb.active
    ws1=wb.create_sheet("mysheet2")
    ws1.append(['이름','국어','영어'])
    for n in myDataList:
        count += 1
        ws1.append([n['name'],n['korean'],n['english']])
 
    #Chart 1
    chart1 = BarChart()
    chart1.style=11
    chart1.title='Bar chart'
    chart1.x_axis.title='이름'
    chart1.y_axis.title='점수'
    data = Reference(ws1, min_col=1, min_row=1, max_row=count+1, max_col=3)
    cat = Reference(ws1, min_col=1, min_row=2, max_row=count+1)
    chart1.add_data(data, titles_from_data=True)
    chart1.set_categories(cat)
    ws1.add_chart(chart1, 'F1')

    wb.save('day2result.xlsx')
    print('write...done')

    
    
def print_menu():
    print("메인메뉴)")
    print("    1.입력")
    print("    2.출력")
    print("    3.엑셀저장(이름별 성적 바차트포함)")
    print("    4.종료")
    return int(input("번호를 입력하세요:"))
    
#    print("메인메뉴","1.입력","2.입력","3.출력","4.종료", sep="\n")

def main():
    while True:
        select = print_menu()
        
        if select == 1:
            input_data()
        elif select == 2:
            print_data()
        elif select == 3:
            save_excel_data()
        elif select == 4:
            break

if __name__ == '__main__':
    main()